from copy import deepcopy
from docx import Document
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SOURCE = 'output/documentation/DocTrackTuao_Final_Thesis_Documentation.docx'
OUT = 'output/documentation/DocTrackTuao_Kevin_Catral_Complete_Final.docx'

d = Document(SOURCE)

def find(text):
    for paragraph in d.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise ValueError(text)

def set_text(paragraph, text):
    paragraph.clear()
    paragraph.add_run(text)

def before(anchor, text='', style=None, page_break=False):
    paragraph = d.add_paragraph(style=style)
    if page_break:
        paragraph.add_run().add_break(WD_BREAK.PAGE)
    if text:
        paragraph.add_run(text)
    anchor._p.addprevious(paragraph._p)
    return paragraph

def before_table(anchor, headers, rows):
    table = d.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value
    anchor._p.addprevious(table._tbl)
    return table

# Match Catral's chapter-label convention and keep every major thesis block on a new page.
chapter_titles = {
    'CHAPTER I THE PROBLEM AND ITS BACKGROUND': 'CHAPTER I. THE PROBLEM AND ITS BACKGROUND',
    'CHAPTER II REVIEW OF RELATED LITERATURE AND STUDIES': 'CHAPTER II. REVIEW OF RELATED LITERATURE AND STUDIES',
    'CHAPTER III RESEARCH METHODOLOGY': 'CHAPTER III. RESEARCH METHODOLOGY',
    'CHAPTER IV RESULTS AND DISCUSSIONS': 'CHAPTER IV. RESULTS AND DISCUSSIONS',
    'CHAPTER V SUMMARY OF FINDINGS AND RECOMMENDATIONS': 'CHAPTER V. SUMMARY OF FINDINGS AND RECOMMENDATIONS',
}
for old, new in chapter_titles.items():
    set_text(find(old), new)

for heading in [
    'CHAPTER II. REVIEW OF RELATED LITERATURE AND STUDIES',
    'CHAPTER III. RESEARCH METHODOLOGY',
    'CHAPTER IV. RESULTS AND DISCUSSIONS',
    'CHAPTER V. SUMMARY OF FINDINGS AND RECOMMENDATIONS',
    'Bibliography', 'APPENDICES', 'Curriculum Vitae',
]:
    before(find(heading), page_break=True)

# Improve the TOC page as a field-based Word TOC. The generated entries inherit Word's TOC styles.
toc = find('TABLE OF CONTENTS')
toc.paragraph_format.space_after = Pt(12)
for paragraph in d.paragraphs:
    if 'Open in Word and update fields to display the Table of Contents.' in paragraph.text:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.0
    if 'Open in Word and update fields to display the List of Tables.' in paragraph.text or 'Open in Word and update fields to display the List of Figures.' in paragraph.text:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(6)

# Catral's Chapter IV includes a dedicated evaluation discussion. Insert the evidence table before Chapter V.
ch5 = find('CHAPTER V. SUMMARY OF FINDINGS AND RECOMMENDATIONS')
before(ch5, '4.3 Extent of Compliance of the Developed System to ISO/IEC 25010 Software Quality Standards', 'Heading 2')
before(ch5, 'The ISO/IEC 25010 questionnaire in Appendix D is the instrument for formal evaluation. The following implementation evidence identifies what evaluators may assess. Weighted means and verbal interpretations must be computed only from accomplished and validated questionnaires.')
before_table(ch5, ['Quality characteristic', 'Implementation evidence available for evaluation'], [
    ['Functional suitability', 'Document registration, QR lookup, routing, receipt, processing history, audit logging, and reports.'],
    ['Performance efficiency', 'Centralized database queries, pagination, and QR-assisted record retrieval.'],
    ['Compatibility', 'Browser-accessible Vue interface and Laravel API integration.'],
    ['Usability', 'Role-aware navigation, clear workflow actions, manual QR entry, image scanning, and responsive forms.'],
    ['Reliability', 'Controlled routing states, receipt confirmation, validation, and preserved histories.'],
    ['Security', 'Authentication, role restrictions, office scope, server-side validation, and audit records.'],
    ['Maintainability', 'Laravel/Vue separation, migrations, controllers, reusable UI components, and automated tests.'],
    ['Portability', 'Web application deployable through supported browser, PHP, MySQL, and Apache environments.'],
])

# Make the Appendix C manual match Catral's fuller task sequence using the already captured system screens.
appendix_d = find('Appendix D ISO 25010 Questionnaire')
for heading, body in [
    ('C.3.1 Routing Follow Up', 'Open the document details page to review its current office, route status, recipient history, and processing information. Follow-up actions must comply with the document custody and pending-route rules displayed by the system.'),
    ('C.3.2 Notification and Reminder Review', 'Use the notification area to review updates relevant to the signed-in user. Notifications support awareness; authorization is still checked by the server before any action is completed.'),
    ('C.4.1 Dashboard and Document Inquiry', 'Use the dashboard for summary monitoring. Use Document Inquiry to search an authorized document record and review its movement and processing history.'),
    ('C.4.2 QR Code Inventory and Scanning', 'Open QR Codes to generate, assign, review, or scan QR references according to the allowed role. A QR reference retrieves a document record but does not itself grant authority to modify the record.'),
    ('C.5.1 Office Directory and Master Data', 'Administrators manage offices and approved master data. Records personnel may view master data where allowed, but cannot use hidden controls to bypass API authorization.'),
    ('C.5.2 User Accounts Roles and Permissions', 'Administrators use User Management to create and maintain accounts, role assignments, and office assignments. Roles determine permitted functions; office assignment determines which currently held documents may be acted on.'),
    ('C.6.1 Basic Troubleshooting', 'If a page fails to load, verify the authorized address, network connection, and account session. If scanning is unavailable on a mobile device, use HTTPS or the supported local testing address, then check browser camera permission. Report persistent issues to the system administrator.'),
]:
    before(appendix_d, heading, 'Heading 5')
    before(appendix_d, body)

# Complete the CV form without inventing personal employment or contact details.
cv = find('Curriculum Vitae')
set_text(cv, 'CURRICULUM VITAE')
before(cv, 'The following curriculum vitae entries are limited to researcher details supplied in the manuscript. Complete personal contact information, education, work experience, training, and certifications from verified records before final binding.')

# Refresh fields when the document is opened in Microsoft Word.
update = OxmlElement('w:updateFields')
update.set(qn('w:val'), 'true')
d.settings.element.append(update)
d.save(OUT)
